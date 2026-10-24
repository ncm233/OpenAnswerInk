import os
import json
import time
from pathlib import Path
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

HISTORY_DIR = Path(os.environ.get("APPDATA", str(Path.home()))) / "OpenAnswerInk" / "history"

class HistoryRecord:
    def __init__(
        self,
        record_id: str,
        timestamp: float,
        record_type: str,  # "exam" or "interview"
        question_text: str,
        answer_text: str,
        reasoning_text: str = "",
        images: Optional[List[str]] = None
    ):
        self.record_id = record_id
        self.timestamp = timestamp
        self.record_type = record_type
        self.question_text = question_text
        self.answer_text = answer_text
        self.reasoning_text = reasoning_text
        self.images = images or []

    def to_dict(self) -> Dict[str, Any]:
        return {
            "record_id": self.record_id,
            "timestamp": self.timestamp,
            "record_type": self.record_type,
            "question_text": self.question_text,
            "answer_text": self.answer_text,
            "reasoning_text": self.reasoning_text,
            "image_count": len(self.images)
        }

class HistoryManager:
    def __init__(self):
        HISTORY_DIR.mkdir(parents=True, exist_ok=True)

    def save_record(
        self,
        record_type: str,
        question_text: str,
        answer_text: str,
        reasoning_text: str = "",
        image_pixmaps: Optional[List[Any]] = None
    ) -> str:
        record_id = time.strftime("%Y%m%d_%H%M%S")
        record_folder = HISTORY_DIR / record_id
        record_folder.mkdir(parents=True, exist_ok=True)

        # 保存元数据与正文
        meta = {
            "record_id": record_id,
            "timestamp": time.time(),
            "time_str": time.strftime("%Y-%m-%d %H:%M:%S"),
            "record_type": record_type,
            "question_text": question_text
        }
        with open(record_folder / "meta.json", "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)

        with open(record_folder / "answer.txt", "w", encoding="utf-8") as f:
            f.write(answer_text)

        if reasoning_text:
            with open(record_folder / "reasoning.txt", "w", encoding="utf-8") as f:
                f.write(reasoning_text)

        # 保存图片
        if image_pixmaps:
            for idx, pix in enumerate(image_pixmaps):
                img_path = record_folder / f"{idx}.png"
                pix.save(str(img_path), "PNG")

        return record_id

    def list_records(self) -> List[Dict[str, Any]]:
        records = []
        if not HISTORY_DIR.exists():
            return records

        for item in sorted(HISTORY_DIR.iterdir(), reverse=True):
            if item.is_dir() and (item / "meta.json").exists():
                try:
                    with open(item / "meta.json", "r", encoding="utf-8") as f:
                        meta = json.load(f)
                    records.append(meta)
                except Exception:
                    pass
        return records

    def export_to_docx(self, output_path: str) -> bool:
        """将全部历史题目与AI解答生成排版整洁的 Word 文档"""
        doc = Document()

        # 标题
        title = doc.add_heading("OpenAnswerInk 智能笔试与面试回顾报告", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_paragraph(f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        records = self.list_records()
        if not records:
            doc.add_paragraph("暂无历史记录。")
            doc.save(output_path)
            return True

        for rec in reversed(records):
            rec_id = rec.get("record_id", "")
            folder = HISTORY_DIR / rec_id
            rec_type = "笔试作答" if rec.get("record_type") == "exam" else "面试问答"
            time_str = rec.get("time_str", "")

            doc.add_heading(f"【{rec_type}】 {time_str}", level=1)

            # 插入截图
            img_files = sorted(folder.glob("*.png"))
            for img in img_files:
                try:
                    doc.add_paragraph(f"题目截图：")
                    doc.add_picture(str(img), width=Inches(5.5))
                except Exception as e:
                    doc.add_paragraph(f"[图片加载失败: {e}]")

            # 插入问题文本
            q_text = rec.get("question_text")
            if q_text:
                doc.add_heading("问题与要求：", level=2)
                doc.add_paragraph(q_text)

            # 插入思考过程
            if (folder / "reasoning.txt").exists():
                with open(folder / "reasoning.txt", "r", encoding="utf-8") as f:
                    reasoning = f.read()
                if reasoning.strip():
                    doc.add_heading("深度思考过程 (Reasoning)：", level=2)
                    p = doc.add_paragraph(reasoning)
                    p.runs[0].font.color.rgb = RGBColor(128, 128, 128)
                    p.runs[0].font.size = Pt(9.5)

            # 插入答案正文
            if (folder / "answer.txt").exists():
                with open(folder / "answer.txt", "r", encoding="utf-8") as f:
                    ans = f.read()
                doc.add_heading("AI 建议回答 / 核心代码：", level=2)
                doc.add_paragraph(ans)

            doc.add_paragraph("-" * 40)

        doc.save(output_path)
        return True

history_manager = HistoryManager()
